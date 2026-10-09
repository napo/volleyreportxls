#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Maurizio Napolitano <maurizio.napolitano@gmail.com>
# SPDX-License-Identifier: AGPL-3.0-or-later

"""Extract the reference fixture (inputs + cached results) from VolleyReportXLS.xlsx.

Development tool only. The application never reads the workbook: this script
produces `src/oracle/volleyreportxls-oracle.json`, which the unit
tests use as an oracle. Every extracted value keeps the address of the source
cell so that a mismatch can be traced back to the original formula.

Usage:  python3 tools/extract-oracle.py [path/to/VolleyReportXLS.xlsx]
Requires: openpyxl
"""

import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "VolleyReportXLS.xlsx"
OUT = ROOT / "src" / "oracle" / "volleyreportxls-oracle.json"

CODE_RE = re.compile(r"^[RBAMPD][#+!\-/=]$")
SUMMARY_LABELS = {
    "Punti": "pointsWon",
    "Errori": "errors",
    "Err. Ric": "receptionErrors",
    "Muri sub": "blockedAttacks",
    "Eff. Ric": "receptionEfficiency",
    "Eff. Att.": "attackEfficiency",
    "V-P": "balance",
}
RATIO_LABELS = {"Efficenza": "efficiency", "Positività": "positivity", "Efficacia": "efficacy"}

# Rilevazione: five tables of 12 player rows each (rows used by the Gioc sheets' VLOOKUPs).
SET_ROWS = {1: (2, 13), 2: (16, 27), 3: (30, 41), 4: (44, 55), 5: (58, 69)}

# Regions (min_row, max_row, min_col, max_col) of each statistics block.
GIOC_REGIONS = {
    "match": (3, 23, "A", "K"),
    1: (2, 23, "M", "X"),
    2: (25, 45, "A", "K"),
    3: (25, 45, "M", "X"),
    4: (47, 67, "A", "K"),
    5: (47, 67, "M", "X"),
}
SQUADRA_REGIONS = {
    "match": (2, 23, "A", "L"),
    1: (24, 45, "A", "L"),
    2: (46, 68, "A", "L"),
    3: (69, 91, "A", "L"),
    4: (92, 114, "A", "L"),
    5: (115, 137, "A", "L"),
}


def plain(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.date().isoformat() if isinstance(v, datetime.datetime) else v.isoformat()
    return v


def cell(ws_values, ref):
    return {"cell": ref, "value": plain(ws_values[ref].value)}


def extract_table(wf, wv, block, r, skill, codes):
    if skill == "P":
        # "P=" is a label with its count on the right.
        c = codes[0][0]
        ref = f"{get_column_letter(c + 1)}{r}"
        block["skills"].setdefault("P", {"counts": {}})["counts"]["P="] = cell(wv, ref)
        return
    table = {"counts": {}, "distribution": {}}
    for c, code in codes:
        col = get_column_letter(c)
        table["counts"][code] = cell(wv, f"{col}{r + 1}")
        if skill != "M":
            table["distribution"][code] = cell(wv, f"{col}{r + 2}")
    last = codes[-1][0]
    if wf.cell(r, last + 1).value == "Tot":
        col = get_column_letter(last + 1)
        table["total"] = cell(wv, f"{col}{r + 1}")
        table["distributionTotal"] = cell(wv, f"{col}{r + 2}")
    # Ratios: label in the column left of the first code, value under the first code.
    first = codes[0][0]
    for rr in range(r + 3, r + 6):
        label = wf.cell(rr, first - 1).value
        if label in RATIO_LABELS:
            table[RATIO_LABELS[label]] = cell(wv, f"{get_column_letter(first)}{rr}")
    if skill == "M":
        table.pop("distribution")
    block["skills"][skill] = table


def extract_block(wf, wv, region):
    """Locate skill tables, ratios and summary values inside a block by their labels."""
    r0, r1, c0, c1 = region
    c0, c1 = column_index_from_string(c0), column_index_from_string(c1)
    block = {"skills": {}, "summary": {}}
    for r in range(r0, r1 + 1):
        row = [(c, wf.cell(r, c).value) for c in range(c0, c1 + 1)]
        codes = [(c, v) for c, v in row if isinstance(v, str) and CODE_RE.match(v)]
        if not codes:
            continue
        groups = {}
        for c, code in codes:
            groups.setdefault(code[0], []).append((c, code))
        for skill, group in groups.items():
            extract_table(wf, wv, block, r, skill, group)
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            label = wf.cell(r, c).value
            if label in SUMMARY_LABELS:
                # The value sits right of its label, except where the layout skips a column.
                offset = 1 if wf.cell(r, c + 1).value is not None else 2
                block["summary"][SUMMARY_LABELS[label]] = cell(wv, f"{get_column_letter(c + offset)}{r}")
    return block


def main():
    wf = openpyxl.load_workbook(SRC)
    wv = openpyxl.load_workbook(SRC, data_only=True)

    ev = wv["Evento"]
    event = {
        "competition": ev["B1"].value,
        "date": plain(ev["B2"].value),
        "venue": ev["B3"].value,
        "team": ev["B4"].value,
        "opponent": ev["B5"].value,
        "scores": [
            {"set": s, "team": plain(ev[f"B{4 + 2 * s}"].value), "opponent": plain(ev[f"B{5 + 2 * s}"].value)}
            for s in range(1, 6)
        ],
    }

    at = wv["Atleti"]
    roster = []
    for r in range(3, 18):
        slot = at[f"A{r}"].value
        if slot and at[f"B{r}"].value is not None:
            roster.append(
                {"slot": slot, "number": plain(at[f"B{r}"].value), "role": at[f"C{r}"].value, "name": at[f"D{r}"].value}
            )

    ril = wv["Rilevazione"]
    rilevazione = []
    for s, (a, b) in SET_ROWS.items():
        lines = []
        for r in range(a, b + 1):
            num = ril.cell(r, 2).value
            cells = [ril.cell(r, c).value for c in range(3, 41)]  # C..AN, as read by the Squadra sheet
            while cells and cells[-1] is None:
                cells.pop()
            if num is None and not cells:
                continue
            lines.append({"row": r, "playerNumber": plain(num), "cells": ["" if v is None else str(v) for v in cells]})
        rilevazione.append({"set": s, "lines": lines})

    # The player sheets are identical copies (only the shirt number changes):
    # slot A is enough as oracle for the per-player logic.
    players = {}
    for p in roster[:1]:
        name = f"Gioc {p['slot']}"
        players[p["slot"]] = {
            "number": p["number"],
            "sheet": name,
            **{
                ("match" if k == "match" else f"set{k}"): extract_block(wf[name], wv[name], reg)
                for k, reg in GIOC_REGIONS.items()
            },
        }

    team = {
        ("match" if k == "match" else f"set{k}"): extract_block(wf["Squadra"], wv["Squadra"], reg)
        for k, reg in SQUADRA_REGIONS.items()
    }
    team["match"]["summary"]["errorsIncludingReception"] = cell(wv["Squadra"], "L11")

    tab = wv["Tabellino"]
    tabellino = {"cells": {}}
    for r in list(range(2, 4)) + list(range(7, 22)) + list(range(23, 30)):
        for c in range(2, 28):
            ref = f"{get_column_letter(c)}{r}"
            v = tab[ref].value
            f = wf["Tabellino"][ref].value
            if isinstance(f, str) and f.startswith("="):
                tabellino["cells"][ref] = plain(v)

    vp = wv["Vinti-Persi Squadra"]
    vinti_persi = {ref: plain(vp[ref].value) for ref in ["B3", "C3", "B4", "C4", "B5", "C5", "B6", "C6", "C7", "B8"]}

    fixture = {
        "source": SRC.name,
        "generatedBy": "tools/extract-oracle.py",
        "event": event,
        "roster": roster,
        "rilevazione": rilevazione,
        "oracle": {"players": players, "team": team, "tabellino": tabellino, "vintiPersi": vinti_persi},
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(fixture, ensure_ascii=False, indent=1) + "\n")
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
